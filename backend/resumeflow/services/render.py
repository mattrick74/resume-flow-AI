"""DocumentRenderer (PRD §8, §9): deterministic DOCX templates for resume and cover letter.

Layout and section order are fixed here; a ``TemplateStyle`` only supplies the look
(fonts, sizes, colors, heading treatment, margins).
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from resumeflow.domain.content import CoverLetterContent, ResumeContent
from resumeflow.domain.templates import BUILTIN_TEMPLATES, DEFAULT_TEMPLATE, TemplateStyle

_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# Kept for callers that list the built-in names.
TEMPLATES = {key: style for key, (_, style) in BUILTIN_TEMPLATES.items()}


def _style(template: TemplateStyle | str) -> TemplateStyle:
    if isinstance(template, TemplateStyle):
        return template
    return TEMPLATES.get(template, TEMPLATES[DEFAULT_TEMPLATE])


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def format_date(value: str | None) -> str:
    if not value:
        return ""
    if value.lower() == "present":
        return "Present"
    if len(value) == 7 and value[4] == "-":
        month = int(value[5:7])
        return f"{_MONTHS[month - 1]} {value[:4]}" if 1 <= month <= 12 else value[:4]
    return value


def date_range(start: str | None, end: str | None) -> str:
    left, right = format_date(start), format_date(end)
    return f"{left} – {right}" if left and right else left or right


def _set_font(element_rpr, font: str) -> None:
    fonts = element_rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), font)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)  # a theme font would override the explicit one


def _base_document(style: TemplateStyle) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)  # US Letter
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(style.margin))
    normal = document.styles["Normal"]
    normal.font.name = style.font
    normal.font.size = Pt(style.size)
    normal.font.color.rgb = _rgb(style.text_color)
    _set_font(normal.element.get_or_add_rPr(), style.font)
    normal.paragraph_format.space_after = Pt(style.paragraph_space)
    normal.paragraph_format.space_before = Pt(0)
    bullets = document.styles["List Bullet"]
    bullets.font.name = style.font
    _set_font(bullets.element.get_or_add_rPr(), style.font)
    return document


def _usable_width(document: Document) -> int:
    section = document.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def _bottom_border(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _heading(document: Document, style: TemplateStyle, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text.upper() if style.heading_caps else text)
    run.bold = style.heading_bold
    run.font.name = style.heading_font
    _set_font(run._r.get_or_add_rPr(), style.heading_font)
    run.font.size = Pt(style.heading_size)
    run.font.color.rgb = _rgb(style.heading_color)
    if style.heading_rule:
        _bottom_border(paragraph, style.rule_color)


def _header_block(document: Document, style: TemplateStyle, name: str, contact_bits: list[str], headline: str | None) -> None:
    align = WD_ALIGN_PARAGRAPH.CENTER if style.name_align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    title = document.add_paragraph()
    title.alignment = align
    title.paragraph_format.space_after = Pt(0)
    run = title.add_run(name)
    run.bold = style.name_bold
    run.font.name = style.heading_font
    _set_font(run._r.get_or_add_rPr(), style.heading_font)
    run.font.size = Pt(style.name_size)
    run.font.color.rgb = _rgb(style.name_color)
    if headline:
        line = document.add_paragraph()
        line.alignment = align
        line.paragraph_format.space_after = Pt(0)
        headline_run = line.add_run(headline)
        headline_run.font.size = Pt(style.size + 1)
    if contact_bits:
        contact = document.add_paragraph()
        contact.alignment = align
        contact_run = contact.add_run("  |  ".join(contact_bits))
        contact_run.font.size = Pt(style.size - 1)
        contact_run.font.color.rgb = _rgb(style.muted_color)


def _bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run(text)


def render_resume(content: ResumeContent, template: TemplateStyle | str = DEFAULT_TEMPLATE) -> bytes:
    style = _style(template)
    document = _base_document(style)
    contact = content.contact
    bits = [b for b in (contact.location, contact.phone, contact.email, *contact.links) if b]
    _header_block(document, style, contact.name, bits, content.headline)

    if content.summary:
        _heading(document, style, "Summary")
        document.add_paragraph(content.summary)

    if content.skills:
        _heading(document, style, "Skills")
        document.add_paragraph(", ".join(content.skills))

    if content.experience:
        _heading(document, style, "Experience")
        width = _usable_width(document)
        for role in content.experience:
            line = document.add_paragraph()
            line.paragraph_format.space_before = Pt(4)
            line.paragraph_format.keep_with_next = True
            line.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            title = line.add_run(role.title)
            title.bold = True
            employer = ", ".join(x for x in (role.employer, role.location) if x)
            if employer:
                line.add_run(f"  ·  {employer}")
            dates = date_range(role.start, role.end)
            if dates:
                date_run = line.add_run(f"\t{dates}")
                date_run.font.color.rgb = _rgb(style.muted_color)
            for bullet in role.bullets:
                _bullet(document, bullet.text)
            if role.technologies:
                tech = document.add_paragraph()
                tech.paragraph_format.space_after = Pt(4)
                tech.add_run("Technologies: ").bold = True
                tech.add_run(", ".join(role.technologies))

    if content.education:
        _heading(document, style, "Education")
        for item in content.education:
            degree = ", ".join(x for x in (item.get("degree"), item.get("field")) if x)
            text = " — ".join(x for x in (degree, item.get("institution")) if x)
            if item.get("end_year"):
                text += f" ({item['end_year']})"
            document.add_paragraph(text)

    if content.certifications:
        _heading(document, style, "Certifications")
        for item in content.certifications:
            text = ", ".join(str(x) for x in (item.get("name"), item.get("issuer"), item.get("year")) if x)
            _bullet(document, text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_cover_letter(content: CoverLetterContent, template: TemplateStyle | str = DEFAULT_TEMPLATE) -> bytes:
    style = _style(template)
    document = _base_document(style)
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    contact = content.contact
    bits = [b for b in ((contact.location, contact.phone, contact.email) if contact else ()) if b]
    _header_block(document, style, content.signature, bits, None)
    document.add_paragraph()
    document.add_paragraph(content.date)
    if content.company or content.job_title:
        recipient = document.add_paragraph()
        if content.company:
            recipient.add_run(content.company).bold = True
        if content.job_title:
            recipient.add_run(("\n" if content.company else "") + f"Re: {content.job_title}")
    document.add_paragraph(content.greeting)
    for paragraph in content.paragraphs:
        body = document.add_paragraph(paragraph.text)
        body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.add_paragraph(content.closing).paragraph_format.space_after = Pt(18)
    document.add_paragraph(content.signature)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
