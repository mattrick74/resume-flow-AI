"""Ten fictional candidates (PRD §16) as DOCX/PDF resumes, plus sample job descriptions.

Used by the test suite and by ``scripts/make-sample-data.mjs`` so the operator can
try the full workflow without real personal data. Every person and company here
is invented.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field


@dataclass
class SampleRole:
    title: str
    employer: str
    location: str
    dates: str
    bullets: list[str]


@dataclass
class SampleCandidate:
    name: str
    email: str
    phone: str
    location: str
    headline: str
    summary: str
    skills: list[str]
    roles: list[SampleRole]
    education: list[str]
    certifications: list[str] = field(default_factory=list)
    authorization: str = "U.S. Citizen"


SAMPLE_CANDIDATES: list[SampleCandidate] = [
    SampleCandidate(
        "Avery Quinlan", "avery.quinlan@example.com", "(555) 201-3344", "Austin, TX",
        "Senior Data Engineer",
        "Data engineer with 9 years building batch and streaming pipelines on AWS.",
        ["Python", "SQL", "Apache Spark", "Apache Airflow", "AWS", "Snowflake", "dbt", "Kafka", "Docker", "Terraform"],
        [
            SampleRole("Senior Data Engineer", "Northwind Analytics Inc", "Austin, TX", "Mar 2020 – Present", [
                "Built Spark and Airflow pipelines processing 4 TB of events per day on AWS",
                "Migrated the reporting warehouse to Snowflake with dbt models, cutting query costs by 35%",
                "Introduced Terraform for data infrastructure across three environments",
            ]),
            SampleRole("Data Engineer", "Bluebonnet Logistics LLC", "Dallas, TX", "Jun 2016 – Feb 2020", [
                "Developed Python ETL jobs loading shipment data into PostgreSQL",
                "Streamed tracking events through Kafka to real-time dashboards",
            ]),
        ],
        ["B.S. Computer Science, University of North Texas, 2016"],
        ["AWS Certified Data Engineer – Associate, 2022"],
    ),
    SampleCandidate(
        "Marisol Venn", "marisol.venn@example.com", "(555) 302-7781", "Seattle, WA",
        "Frontend Engineer",
        "Frontend engineer focused on accessible React and TypeScript applications.",
        ["React", "TypeScript", "JavaScript", "Next.js", "CSS", "Jest", "Figma", "GraphQL"],
        [
            SampleRole("Frontend Engineer", "Cascade Commerce Corp", "Seattle, WA", "Jan 2021 – Present", [
                "Led the rebuild of the checkout flow in React and TypeScript, raising conversion by 12%",
                "Created a shared component library used by six product teams",
                "Raised accessibility audit scores to WCAG 2.1 AA across the storefront",
            ]),
            SampleRole("Web Developer", "Rainier Digital Agency", "Tacoma, WA", "Aug 2018 – Dec 2020", [
                "Built marketing sites in Next.js with a GraphQL CMS",
                "Wrote Jest test suites that reduced front-end regressions",
            ]),
        ],
        ["B.A. Interactive Media, Evergreen State College, 2018"],
    ),
    SampleCandidate(
        "Desmond Okafor", "desmond.okafor@example.com", "(555) 410-9902", "Chicago, IL",
        "DevOps Engineer",
        "DevOps engineer running Kubernetes platforms and CI/CD for regulated industries.",
        ["Kubernetes", "Docker", "Terraform", "AWS", "Jenkins", "Python", "Linux", "Prometheus", "Grafana", "Helm"],
        [
            SampleRole("DevOps Engineer", "Lakeshore Health Systems", "Chicago, IL", "Apr 2019 – Present", [
                "Operated 14 Kubernetes clusters with Helm and Terraform on AWS",
                "Rebuilt Jenkins CI/CD pipelines, cutting release time from 2 days to 3 hours",
                "Set up Prometheus and Grafana monitoring with on-call alerting",
            ]),
            SampleRole("Systems Administrator", "Midway Freight Co", "Chicago, IL", "Jul 2015 – Mar 2019", [
                "Administered 200 Linux servers and automated patching with Python",
            ]),
        ],
        ["B.S. Information Technology, Illinois Institute of Technology, 2015"],
        ["Certified Kubernetes Administrator (CKA), 2021"],
    ),
    SampleCandidate(
        "Priya Raman", "priya.raman@example.com", "(555) 515-6620", "Boston, MA",
        "Machine Learning Engineer",
        "ML engineer shipping NLP models to production with PyTorch and Python.",
        ["Python", "PyTorch", "Machine Learning", "Natural Language Processing", "SQL", "Docker", "AWS", "scikit-learn"],
        [
            SampleRole("Machine Learning Engineer", "Harborview AI Labs", "Boston, MA", "Sep 2020 – Present", [
                "Deployed PyTorch NLP models serving 2 million predictions per day",
                "Built evaluation pipelines in Python that caught model regressions before release",
            ]),
            SampleRole("Data Scientist", "Commonwealth Insurance Group", "Boston, MA", "Jun 2017 – Aug 2020", [
                "Developed scikit-learn claim-risk models adopted by the underwriting team",
                "Wrote SQL feature pipelines over 50 million policy records",
            ]),
        ],
        ["M.S. Computer Science, Northeastern University, 2017"],
        authorization="Permanent Resident (Green Card)",
    ),
    SampleCandidate(
        "Theo Lindqvist", "theo.lindqvist@example.com", "(555) 620-1187", "Denver, CO",
        "Backend Engineer",
        "Backend engineer building Java and Spring Boot microservices.",
        ["Java", "Spring Boot", "Microservices", "PostgreSQL", "Kafka", "Docker", "Kubernetes", "REST APIs"],
        [
            SampleRole("Senior Backend Engineer", "Front Range Payments Inc", "Denver, CO", "Feb 2019 – Present", [
                "Designed Spring Boot microservices handling 3,000 payment requests per second",
                "Moved services to Kubernetes with zero-downtime deployments",
                "Mentored four engineers on API design and code review",
            ]),
            SampleRole("Software Engineer", "Aspen Software Solutions", "Boulder, CO", "May 2014 – Jan 2019", [
                "Built REST APIs in Java backed by PostgreSQL",
            ]),
        ],
        ["B.S. Software Engineering, Colorado School of Mines, 2014"],
    ),
    SampleCandidate(
        "Nadia Ferreira", "nadia.ferreira@example.com", "(555) 731-4456", "Miami, FL",
        "Business Analyst",
        "Business analyst translating stakeholder needs into requirements for finance systems.",
        ["Business Analysis", "Requirements Gathering", "SQL", "Microsoft Excel", "Power BI", "Jira", "Agile", "Stakeholder Management"],
        [
            SampleRole("Senior Business Analyst", "Biscayne Financial Partners", "Miami, FL", "Jan 2019 – Present", [
                "Gathered requirements for a loan-origination platform used by 300 staff",
                "Built Power BI dashboards that replaced 20 manual Excel reports",
            ]),
            SampleRole("Business Analyst", "Coral Gables Insurance Co", "Coral Gables, FL", "Jun 2015 – Dec 2018", [
                "Wrote user stories and acceptance criteria in Jira for Agile teams",
            ]),
        ],
        ["B.B.A. Information Systems, Florida International University, 2015"],
        ["Certified Business Analysis Professional (CBAP), 2020"],
    ),
    SampleCandidate(
        "Gideon Marsh", "gideon.marsh@example.com", "(555) 842-2093", "Atlanta, GA",
        "Cloud Security Engineer",
        "Security engineer hardening AWS and Azure environments.",
        ["Cybersecurity", "AWS", "Microsoft Azure", "Terraform", "Python", "Splunk", "Linux", "Networking"],
        [
            SampleRole("Cloud Security Engineer", "Peachtree Secure Cloud LLC", "Atlanta, GA", "Mar 2018 – Present", [
                "Automated AWS and Azure security baselines with Terraform and Python",
                "Built Splunk detections that reduced incident response time by 40%",
            ]),
            SampleRole("Network Engineer", "Southern Telecom Corp", "Atlanta, GA", "Aug 2013 – Feb 2018", [
                "Managed enterprise networking for 12 regional offices",
            ]),
        ],
        ["B.S. Cybersecurity, Georgia State University, 2013"],
        ["CISSP, 2019", "AWS Certified Security – Specialty, 2021"],
        authorization="U.S. Citizen; active Secret clearance",
    ),
    SampleCandidate(
        "Lena Hoffmann", "lena.hoffmann@example.com", "(555) 953-6718", "New York, NY",
        "Product Manager",
        "Product manager for B2B SaaS analytics products.",
        ["Product Management", "Agile", "Scrum", "Jira", "SQL", "Stakeholder Management", "Data Analysis"],
        [
            SampleRole("Senior Product Manager", "Hudson Metrics Inc", "New York, NY", "Apr 2020 – Present", [
                "Owned the analytics product roadmap, growing annual revenue from $4M to $9M",
                "Ran discovery interviews with 60 customers to prioritise features",
            ]),
            SampleRole("Product Manager", "Brooklyn Data Co", "Brooklyn, NY", "Jan 2016 – Mar 2020", [
                "Launched a self-serve reporting tool adopted by 1,200 accounts",
            ]),
        ],
        ["MBA, Columbia Business School, 2015"],
        authorization="Requires H-1B visa sponsorship",
    ),
    SampleCandidate(
        "Rafael Ortiz", "rafael.ortiz@example.com", "(555) 164-3390", "Phoenix, AZ",
        "QA Automation Engineer",
        "QA automation engineer building Selenium and Playwright test frameworks.",
        ["Selenium", "Playwright", "Python", "Java", "Jenkins", "Quality Assurance", "REST APIs", "Jira"],
        [
            SampleRole("QA Automation Engineer", "Sonoran Software Inc", "Phoenix, AZ", "May 2019 – Present", [
                "Built a Playwright regression suite covering 800 end-to-end scenarios",
                "Integrated automated tests into Jenkins, blocking 90% of regressions before release",
            ]),
            SampleRole("QA Analyst", "Mesa Health Technologies", "Mesa, AZ", "Jan 2016 – Apr 2019", [
                "Automated REST API tests in Python and Selenium UI tests in Java",
            ]),
        ],
        ["B.S. Computer Information Systems, Arizona State University, 2015"],
    ),
    SampleCandidate(
        "Hana Kobayashi", "hana.kobayashi@example.com", "(555) 275-8841", "San Jose, CA",
        "Full Stack Engineer",
        "Full stack engineer working across Python, React and PostgreSQL.",
        ["Python", "Django", "React", "TypeScript", "PostgreSQL", "AWS", "Docker", "REST APIs"],
        [
            SampleRole("Full Stack Engineer", "Silicon Orchard Labs", "San Jose, CA", "Jul 2021 – Present", [
                "Built Django REST APIs and React dashboards for a logistics platform",
                "Reduced PostgreSQL query latency by 60% through indexing and query tuning",
            ]),
            SampleRole("Junior Developer", "Bay Area Web Studio", "Oakland, CA", "Aug 2019 – Jun 2021", [
                "Maintained React front ends and Python back ends for client projects",
            ]),
        ],
        ["B.S. Computer Engineering, San Jose State University, 2019"],
    ),
]

SAMPLE_JOBS: list[dict[str, str]] = [
    {
        "url": "https://careers.example.com/jobs/senior-data-engineer-4821?utm_source=board",
        "company": "Acme Data Co",
        "title": "Senior Data Engineer",
        "text": """Senior Data Engineer
Acme Data Co - Austin, TX (Hybrid)

About the role
You will design and run the data pipelines behind our analytics products.

Responsibilities
- Build batch and streaming pipelines with Spark and Airflow
- Model warehouse data in Snowflake using dbt
- Own data infrastructure as code with Terraform

Requirements
- 5+ years of data engineering experience
- Strong Python and SQL
- Experience with Apache Spark, Apache Airflow and AWS
- Must be authorized to work in the United States without sponsorship

Preferred
- Kafka
- dbt
- Snowflake
""",
    },
    {
        "url": "https://jobs.example.org/contoso/frontend-engineer",
        "company": "Contoso Retail",
        "title": "Frontend Engineer",
        "text": """Frontend Engineer
Contoso Retail - Remote (US)

What you'll do
- Build accessible customer-facing experiences in React and TypeScript
- Collaborate with designers in Figma

What you bring
- 3+ years building React applications
- TypeScript, JavaScript, CSS
- Testing with Jest

Nice to have
- Next.js
- GraphQL
""",
    },
]


def resume_docx_bytes(candidate: SampleCandidate) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    heading = document.add_paragraph()
    run = heading.add_run(candidate.name)
    run.bold = True
    run.font.size = Pt(18)
    document.add_paragraph(candidate.headline)
    document.add_paragraph(f"{candidate.location} | {candidate.phone} | {candidate.email}")
    document.add_paragraph("Summary")
    document.add_paragraph(candidate.summary)
    document.add_paragraph("Experience")
    for role in candidate.roles:
        document.add_paragraph(f"{role.title} | {role.employer} | {role.location} | {role.dates}")
        for bullet in role.bullets:
            document.add_paragraph(f"• {bullet}")
    document.add_paragraph("Skills")
    document.add_paragraph(", ".join(candidate.skills))
    document.add_paragraph("Education")
    for line in candidate.education:
        document.add_paragraph(line)
    if candidate.certifications:
        document.add_paragraph("Certifications")
        for line in candidate.certifications:
            document.add_paragraph(line)
    document.add_paragraph("Work Authorization")
    document.add_paragraph(candidate.authorization)
    buffer = io.BytesIO()
    document.save(buffer)
    return _without_timestamps(buffer.getvalue())


def _without_timestamps(docx: bytes) -> bytes:
    """Re-packs a DOCX with fixed zip entry times so the same sample always has the same bytes (and hash)."""
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(fixed, source.read(info.filename))
    return out.getvalue()


def resume_pdf_bytes(candidate: SampleCandidate) -> bytes:
    import pymupdf

    lines = [
        candidate.name,
        candidate.headline,
        f"{candidate.location} | {candidate.phone} | {candidate.email}",
        "",
        "Summary",
        candidate.summary,
        "",
        "Experience",
    ]
    for role in candidate.roles:
        lines.append(f"{role.title} | {role.employer} | {role.location} | {role.dates}")
        lines.extend(f"- {b}" for b in role.bullets)
    lines += ["", "Skills", ", ".join(candidate.skills), "", "Education", *candidate.education]
    if candidate.certifications:
        lines += ["", "Certifications", *candidate.certifications]
    lines += ["", "Work Authorization", candidate.authorization]
    # The built-in PDF font has no en dash; real resumes keep theirs.
    lines = [line.replace("–", "-") for line in lines]
    document = pymupdf.open()
    page = document.new_page()
    y = 50
    for line in lines:
        if y > 800:
            page = document.new_page()
            y = 50
        page.insert_text((50, y), line, fontsize=9)
        y += 13
    data = document.tobytes()
    document.close()
    return data
