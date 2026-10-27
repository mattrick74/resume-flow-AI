"""Contact links: a hyperlinked label keeps its address at import, a bare label never reaches a resume, and the resume
shows the address itself as a clickable link."""

from __future__ import annotations

import io

import pymupdf
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from resumeflow.devtools.sample_data import SAMPLE_CANDIDATES, resume_docx_bytes, resume_pdf_bytes
from resumeflow.domain.content import Contact, ResumeContent
from resumeflow.domain.links import clean_links, link_text
from resumeflow.services.render import render_resume
from resumeflow.services.resume_parser import extract_text, parse_resume

PROFILE = "https://www.linkedin.com/in/avery-quinlan/"


def _docx_with_hyperlinked_label() -> bytes:
    """A sample resume whose contact line ends with "LinkedIn", hyperlinked to the profile (as Word saves it)."""
    document = Document(io.BytesIO(resume_docx_bytes(SAMPLE_CANDIDATES[0])))
    contact = document.paragraphs[2]  # location | phone | email
    contact.add_run(" | ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), contact.part.relate_to(PROFILE, RT.HYPERLINK, is_external=True))
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "LinkedIn"
    run.append(text)
    link.append(run)
    contact._p.append(link)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_with_hyperlinked_label() -> bytes:
    """The sample resume as a PDF whose contact line ends with "LinkedIn", linked to the profile."""
    original = pymupdf.open(stream=resume_pdf_bytes(SAMPLE_CANDIDATES[0]), filetype="pdf")
    lines = original[0].get_text("text").splitlines()
    original.close()
    lines = [f"{line} | LinkedIn" if "avery.quinlan@example.com" in line else line for line in lines]
    document = pymupdf.open()
    page = document.new_page()
    for index, line in enumerate(lines):
        page.insert_text((50, 50 + 13 * index), line, fontsize=9)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": page.search_for("LinkedIn")[0], "uri": PROFILE})
    return document.tobytes()


def test_a_hyperlinked_label_keeps_its_address_at_import():
    for fmt, data in (("docx", _docx_with_hyperlinked_label()), ("pdf", _pdf_with_hyperlinked_label())):
        text = extract_text(fmt, data)
        assert f"LinkedIn ({PROFILE})" in text, fmt
        assert parse_resume(text).profile.links == [PROFILE], fmt


def test_only_real_addresses_are_links():
    assert clean_links(["LinkedIn", "GitHub", "", "jane@example.com"]) == []
    assert clean_links(["LinkedIn: linkedin.com/in/jane", "https://www.linkedin.com/in/jane/", "github.com/jane"]) == [
        "linkedin.com/in/jane",
        "github.com/jane",
    ]
    assert link_text("https://www.linkedin.com/in/jane/") == "linkedin.com/in/jane"


def test_the_resume_shows_the_address_as_a_clickable_link():
    content = ResumeContent(
        contact=Contact(name="Avery Quinlan", email="avery@example.com", location="Austin, TX", links=["LinkedIn", PROFILE]),
        summary="Data engineer.",
    )
    document = Document(io.BytesIO(render_resume(content)))
    contact = next(p for p in document.paragraphs if "avery@example.com" in p._p.xml)
    shown = "".join(t.text or "" for t in contact._p.iter(qn("w:t")))
    assert shown == "Austin, TX  |  avery@example.com  |  linkedin.com/in/avery-quinlan"  # no bare "LinkedIn"
    links = list(contact._p.iter(qn("w:hyperlink")))
    assert len(links) == 1
    assert contact.part.rels[links[0].get(qn("r:id"))].target_ref == PROFILE
